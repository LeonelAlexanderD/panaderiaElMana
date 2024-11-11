import csv
from datetime import datetime
from decimal import Decimal
import json
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Sum
import openpyxl
from xhtml2pdf import pisa
from django.template.loader import render_to_string
from django.db.models import F
from django.utils.dateparse import parse_date


from productos.models import Insumo, Producto
from usuarios.decorators import perfil_gerente_o_superior
from usuarios.models import Empleado
from ventas.forms import ClienteForm
from ventas.models import CarritoProducto, Cliente_Mayorista, Comprobante, Venta

# Create your views here.
# 

@login_required(login_url='usuarios:login')
def nueva_venta(request):
    productos = Producto.objects.order_by('categoria', 'subcategoria')
    productos_dict = {}
    for producto in productos:
        if producto.categoria not in productos_dict:
            productos_dict[producto.categoria] = {}
        if producto.subcategoria not in productos_dict[producto.categoria]:
            productos_dict[producto.categoria][producto.subcategoria] = []
        productos_dict[producto.categoria][producto.subcategoria].append(producto)
    
    tipo_venta_choices = dict(Comprobante.TIPO_VENTA)
    forma_pago_choices = dict(Comprobante.FORMA_DE_PAGO)
    tipo_comprobante_choices = dict(Comprobante.TIPO_COMPROBANTE)
        
    if 'comprobante_temp' not in request.session:
        request.session['comprobante_temp'] = {
            'tipo_de_venta': None,
            'forma_de_pago': None,
            'tipo_comprobante': None,
            'observacion': '',
            'items': []
        }
    
    return render(request, 'venta/nueva_venta.html',{
        'productos': productos_dict, 
        'comprobante_temp': request.session['comprobante_temp'],
        'tipo_venta_choices': json.dumps(tipo_venta_choices),
        'forma_pago_choices': json.dumps(forma_pago_choices),
        'tipo_comprobante_choices': json.dumps(tipo_comprobante_choices),
        })


@login_required(login_url='usuarios:login')
def agregar_producto_carrito(request):
    if request.method == 'POST':
        producto_id = request.POST.get('producto_id')
        cantidad = Decimal(request.POST.get('cantidad', 1))
        producto = Producto.objects.get(id=producto_id)
        
        # Crear o actualizar el comprobante temporal en la sesión
        # if 'comprobante_temp' not in request.session:
        #     crear_comprobante_temp(request)
        #inicialmente tengo la vista crear_comprobante_temp donde creo un comprobante temporal, pero 
        #movi la logica hacia listar productos, tengo que crear nuevamente la vista para el comprobante
        #temporal si veo que necesito manipular 
        
        comprobante_temp = request.session['comprobante_temp']
        item_existente = next((item for item in comprobante_temp['items'] if item['producto_id'] == producto.id), None)
        
        if item_existente:
            # # Si el producto ya está en el carrito, actualiza la cantidad
            item_existente['cantidad'] += float(cantidad)
            item_existente['subtotal'] = item_existente['cantidad'] * float(producto.precio)
        else:
            ## Si no está en el carrito, agregarlo como nuevo
            subtotal = producto.precio * cantidad
            comprobante_temp['items'].append({
                'imagen': producto.imagen.url,
                'producto_id': producto.id,
                'nombre': producto.nombre,
                'cantidad': float(cantidad),
                'precio': float(producto.precio),
                'subtotal': float(subtotal)
            })
        
        request.session['comprobante_temp'] = comprobante_temp
        return redirect('ventas:nueva_venta') 


@login_required(login_url='usuarios:login')
def actualizar_o_eliminar_producto(request):
    if request.method == 'POST':
        producto_id = int(request.POST.get('producto_id'))
        accion = request.POST.get('accion')
        comprobante_temp = request.session.get('comprobante_temp', {})

        item = next((item for item in comprobante_temp['items'] if item['producto_id'] == producto_id), None)

        if item:
            if accion == 'actualizar':
                nueva_cantidad = int(request.POST.get('cantidad', 1))
                item['cantidad'] = nueva_cantidad  
                item['subtotal'] = nueva_cantidad * item['precio']  
            elif accion == 'eliminar':
                comprobante_temp['items'].remove(item)

            request.session['comprobante_temp'] = comprobante_temp
            request.session.modified = True
            
            total_carrito = sum(i['subtotal'] for i in comprobante_temp['items'])
            return JsonResponse({'success': True, 'total_carrito': total_carrito, 'subtotal': item['subtotal'] if accion == 'actualizar' else 0})

    return JsonResponse({'success': False}, status=400)
        

##    
@login_required(login_url='usuarios:login')
def generar_comprobante(request):
    if request.method == 'POST':        
        tipo_venta = request.POST.get('tipo_de_venta')
        forma_pago = request.POST.get('forma_de_pago')
        tipo_comprobante = request.POST.get('tipo_comprobante')
        observacion = request.POST.get('observacion')
        
        if not all([tipo_venta, forma_pago, tipo_comprobante]):
            return redirect('ventas:nueva_venta')
        
        ## creando comprobante
        comprobante = Comprobante(
            tipo_de_venta = tipo_venta,
            forma_de_pago = forma_pago,
            tipo_comprobante = tipo_comprobante,
            total_comprobante = 0,
            observacion = observacion
        )
        comprobante.save()
        
        for item in request.session['comprobante_temp']['items']:
            producto = Producto.objects.get(id=item['producto_id'])
            cantidad = item['cantidad']
            carrito_producto = CarritoProducto.objects.create(
                comprobante = comprobante,
                producto = producto,
                cantidad = cantidad
            )
            producto.stock -= Decimal(cantidad)
            producto.save()
            
        
        empleado = Empleado.objects.get(usuario=request.user)
        venta = Venta(comprobante = comprobante, vendedor = empleado)
        venta.save()
        
        del request.session['comprobante_temp']
        
        return redirect('ventas:ver_comprobante', comprobante_id=comprobante.id)
    
    return redirect('ventas:nueva_venta')




@login_required(login_url='usuarios:login')
def ver_comprobante(request, comprobante_id):
    comprobante = Comprobante.objects.get(id=comprobante_id)
    items = comprobante.items.all()
    # total_comprobante = sum(item.subtotal for item in items)
    comprobante.actualizarTotalComprobante()
    
    return render(request, 'venta/comprobante.html', {
        'comprobante': comprobante,
        'items': items,
        # 'total_comprobante': total_comprobante,
        })




@login_required(login_url='usuarios:login')
def ver_detalles_venta(request, id):
    venta = get_object_or_404(Venta, id=id)
    empleado = venta.vendedor
    comprobante = venta.comprobante
    contexto = {
        'empleado': empleado,
        'comprobante': comprobante,
        'items': comprobante.items.all(),
    }
    return render(request, 'venta/detalles_venta.html', contexto)


@login_required(login_url='usuarios:login')
def listar_ventas(request):
    ventas = Venta.objects.all()
    return render(request, 'venta/lista_ventas.html', {'ventas': ventas})


@login_required(login_url='usuarios:login')
def registrar_cliente(request):   
    if request.method == "POST":
        form = ClienteForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('ventas:listar_clientes')
    else:
        form = ClienteForm()
    return render(request, 'gestion/lista_clientes.html', {'form':form})

@login_required(login_url='usuarios:login')
def listar_clientes(request):        
    clientes = Cliente_Mayorista.objects.all()
    form = ClienteForm()
    print(clientes)
    return render(request,'gestion/lista_clientes.html', {'clientes': clientes, 'form':form})



#########################################################
####    informes
@perfil_gerente_o_superior
def informes(request):
    return render(request,'gestion/informes.html')

#productos tabla
@perfil_gerente_o_superior
def productos_mas_vendidos(request):
    productos_venta = (
        Venta.objects
        .values('comprobante__items__producto__nombre')
        .annotate(total_vendido=Sum('comprobante__items__cantidad'))
        .order_by('-total_vendido')
    )
    
    return render(request, 'gestion/productos_mas_vendidos.html', {'productos_venta': productos_venta})


## csv
@perfil_gerente_o_superior
def exportar_csv(request):
    productos_venta = (
        Venta.objects
        .values('comprobante__items__producto__nombre')
        .annotate(total_vendido=Sum('comprobante__items__cantidad'))
        .order_by('-total_vendido')
    )

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="productos_mas_vendidos.csv"'

    writer = csv.writer(response)
    writer.writerow(['Producto', 'Cantidad Vendida'])

    for producto in productos_venta:
        writer.writerow([producto['comprobante__items__producto__nombre'], producto['total_vendido']])

    return response

##excel
@perfil_gerente_o_superior
def exportar_excel(request):
    productos_venta = (
        Venta.objects
        .values('comprobante__items__producto__nombre')
        .annotate(total_vendido=Sum('comprobante__items__cantidad'))
        .order_by('-total_vendido')
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Productos Más Vendidos"

    ws.append(['Producto', 'Cantidad Vendida'])

    for producto in productos_venta:
        ws.append([producto['comprobante__items__producto__nombre'], producto['total_vendido']])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="productos_mas_vendidos.xlsx"'

    wb.save(response)
    return response

##pdf
@perfil_gerente_o_superior
def exportar_pdf(request):
    productos_venta = (
        Venta.objects
        .values('comprobante__items__producto__nombre')
        .annotate(total_vendido=Sum('comprobante__items__cantidad'))
        .order_by('-total_vendido')
    )

    context = {'productos_venta': productos_venta}
    html = render_to_string('gestion/productos_mas_vendidos.html', context)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="productos_mas_vendidos.pdf"'

    pisa_status = pisa.CreatePDF(html, dest=response)

    if pisa_status.err:
        return HttpResponse('Error generando el PDF', status=500)

    return response



##materia prima
@perfil_gerente_o_superior
def insumos_faltantes(request):
    insumos_bajo_stock = Insumo.objects.filter(stock__lte=F('punto_de_pedido'))
    return render(request, 'gestion/lista_insumos_faltantes.html',{'insumos':insumos_bajo_stock})


#pdf
@perfil_gerente_o_superior
def exportar_materia_faltante_pdf(request):
    insumos = Insumo.objects.filter(stock__lte=F('punto_de_pedido'))

    context = {'insumos': insumos}
    html = render_to_string('gestion/lista_insumos_faltantes.html', context)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="lista_insumos_faltantes.pdf"'

    pisa_status = pisa.CreatePDF(html, dest=response)

    if pisa_status.err:
        return HttpResponse('Error generando el PDF', status=500)

    return response

#csv
@perfil_gerente_o_superior
def exportar_materia_faltante_csv(request):
    insumos_bajo_stock = Insumo.objects.filter(stock__lte=F('punto_de_pedido'))

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="materia_prima_faltante.csv"'

    writer = csv.writer(response)
    writer.writerow(['Insumo', 'Stock Actual', 'Punto de Pedido'])

    for insumo in insumos_bajo_stock:
        writer.writerow([insumo.nombre, insumo.stock, insumo.punto_de_pedido])

    return response

#excel xlsx
@perfil_gerente_o_superior
def exportar_materia_faltante_xlsx(request):
    insumos_bajo_stock = Insumo.objects.filter(stock__lte=F('punto_de_pedido'))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Materia Prima Faltante'
    ws.append(['Insumo', 'Stock Actual', 'Punto de Pedido'])

    for insumo in insumos_bajo_stock:
        ws.append([insumo.nombre, insumo.stock, insumo.punto_de_pedido])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="materia_prima_faltante.xlsx"'

    wb.save(response)

    return response



## ventas reporte
@perfil_gerente_o_superior
def reporte_ventas(request):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')

    ventas = Venta.objects.all()
    
    if fecha_inicio and fecha_fin:
        ventas = ventas.filter(
            fecha__range=[parse_date(fecha_inicio), parse_date(fecha_fin)]
        )

    context = {
        'ventas': ventas,
        'fecha_inicio': fecha_inicio,
        'fecha_fin': fecha_fin,
    }
    return render(request, 'gestion/listado_ventas.html', context)


#pdf
@perfil_gerente_o_superior
def exportar_ventas_pdf(request):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')

    ventas = Venta.objects.all()
    if fecha_inicio and fecha_fin:
        ventas = ventas.filter(
            fecha__range=[parse_date(fecha_inicio), parse_date(fecha_fin)]
        )

    context = {'ventas': ventas}
    html = render_to_string('gestion/listado_ventas.html', context)
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="listado_ventas.pdf"'

    pisa_status = pisa.CreatePDF(html, dest=response)
    if pisa_status.err:
        return HttpResponse('Error al generar PDF', status=500)
    
    return response


#csv
@perfil_gerente_o_superior
def exportar_ventas_csv(request):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')

    ventas = Venta.objects.all()
    if fecha_inicio and fecha_fin:
        ventas = ventas.filter(
            fecha__range=[parse_date(fecha_inicio), parse_date(fecha_fin)]
        )

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="listado_ventas.csv"'

    writer = csv.writer(response)
    writer.writerow(['Fecha', 'Vendedor', 'Total de la Venta'])

    for venta in ventas:
        fecha_formateada = venta.fecha.strftime('%Y-%m-%d')
        writer.writerow([fecha_formateada, f"{venta.vendedor.nombre} {venta.vendedor.apellido}", venta.comprobante.total_comprobante])

    return response

#xlsx
@perfil_gerente_o_superior
def exportar_ventas_xlsx(request):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')

    if fecha_inicio and fecha_fin:
        fecha_inicio = datetime.strptime(fecha_inicio, '%Y-%m-%d')
        fecha_fin = datetime.strptime(fecha_fin, '%Y-%m-%d')

        ventas = Venta.objects.filter(fecha__range=[fecha_inicio, fecha_fin]).select_related('vendedor')
    else:
        ventas = Venta.objects.all().select_related('vendedor')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventas"

    ws.append(['Fecha', 'Vendedor', 'Total de la Venta'])

    for venta in ventas:
        vendedor_nombre = f"{venta.vendedor.nombre} {venta.vendedor.apellido}"
        ws.append([venta.fecha.strftime('%Y-%m-%d'), vendedor_nombre, venta.comprobante.total_comprobante])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="ventas.xlsx"'

    wb.save(response)
    return response